"""Binary resume/cover-letter artifact rendering.

The source content remains deterministic plain text; this module renders the
same approved content into uploadable PDF and DOCX artifacts without changing
the underlying claims.
"""

from __future__ import annotations

import hashlib
import io
import re
from typing import Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def _safe_filename(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip())
    return value.strip("-") or "document"


def _render_pdf(content: str) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=(8.5 * inch, 11 * inch),
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
        title="AI Job Agent document",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ArtifactTitle", parent=styles["Title"], alignment=TA_CENTER, fontSize=16,
        leading=19, spaceAfter=8,
    )
    heading_style = ParagraphStyle(
        "ArtifactHeading", parent=styles["Heading2"], fontSize=10.5,
        leading=13, spaceBefore=7, spaceAfter=3,
    )
    body_style = ParagraphStyle(
        "ArtifactBody", parent=styles["BodyText"], fontSize=9.5,
        leading=12, spaceAfter=4,
    )

    lines = [line.rstrip() for line in content.splitlines()]
    story = []
    first_nonempty = True
    headings = {"SUMMARY", "SKILLS", "EXPERIENCE", "PROJECTS", "EDUCATION"}

    for line in lines:
        if not line:
            story.append(Spacer(1, 2))
            continue
        escaped = (
            line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        )
        if first_nonempty:
            story.append(Paragraph(escaped, title_style))
            first_nonempty = False
        elif line.strip().upper() in headings:
            story.append(Paragraph(escaped, heading_style))
        else:
            story.append(Paragraph(escaped, body_style))

    doc.build(story)
    return out.getvalue()


def _render_docx(content: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = 0.55 * inch
    section.bottom_margin = 0.55 * inch
    section.left_margin = 0.65 * inch
    section.right_margin = 0.65 * inch

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = __import__("docx").shared.Pt(10)

    headings = {"SUMMARY", "SKILLS", "EXPERIENCE", "PROJECTS", "EDUCATION"}
    first_nonempty = True
    for line in content.splitlines():
        if not line.strip():
            continue
        p = doc.add_paragraph()
        if first_nonempty:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line.strip())
            run.bold = True
            run.font.size = __import__("docx").shared.Pt(16)
            first_nonempty = False
        elif line.strip().upper() in headings:
            run = p.add_run(line.strip())
            run.bold = True
            run.font.size = __import__("docx").shared.Pt(11)
        else:
            p.add_run(line.strip())

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def build_artifacts(
    content: str,
    *,
    doc_type: str,
    user_id: int,
    job_id: int,
    version: int,
) -> Dict[str, dict]:
    """Render PDF and DOCX artifacts and return bytes plus auditable metadata."""
    stem = _safe_filename(f"{doc_type}-user{user_id}-job{job_id}-v{version}")
    rendered = {
        "pdf": (_render_pdf(content), "application/pdf", f"{stem}.pdf"),
        "docx": (
            _render_docx(content),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{stem}.docx",
        ),
    }
    artifacts: Dict[str, dict] = {}
    for fmt, (payload, content_type, filename) in rendered.items():
        artifacts[fmt] = {
            "filename": filename,
            "content_type": content_type,
            "size_bytes": len(payload),
            "sha256": hashlib.sha256(payload).hexdigest(),
            "bytes": payload,
        }
    return artifacts
