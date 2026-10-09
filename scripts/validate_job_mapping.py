#!/usr/bin/env python3
"""Validate the evidence contract for synthetic-to-persisted benchmark job mapping.

This validates completeness and structure; it does not independently authenticate
the source database. The exported source hash and mapping evidence must be reviewed.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_COLUMNS = {
    "benchmark_job_id",
    "persisted_job_id",
    "persisted_job_title",
    "persisted_company",
    "persisted_url",
    "source_table",
    "verified_at_utc",
    "source_export_sha256",
}


def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def read_benchmark_job_ids(path: Path) -> set[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    preferred = ("Final Gold 500", "Benchmark Data")
    sheet = next((workbook[name] for name in preferred if name in workbook.sheetnames), workbook.active)
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        raise ValueError("benchmark workbook has no header row")
    headers = {norm(value): index for index, value in enumerate(header)}
    job_index = headers.get("jobid")
    if job_index is None:
        raise ValueError("benchmark workbook is missing a Job ID column")
    job_ids = {
        str(row[job_index]).strip()
        for row in rows
        if len(row) > job_index and row[job_index] is not None and str(row[job_index]).strip()
    }
    if not job_ids:
        raise ValueError("benchmark workbook contains no job IDs")
    return job_ids


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--mapping", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    args = parser.parse_args()

    if not args.workbook.is_file() or not args.mapping.is_file():
        raise SystemExit("workbook and mapping files must both exist")

    try:
        expected_ids = read_benchmark_job_ids(args.workbook)
        with args.mapping.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            headers = {norm(name): name for name in (reader.fieldnames or [])}
            missing = REQUIRED_COLUMNS - set(headers)
            if missing:
                raise ValueError("mapping columns missing: " + ", ".join(sorted(missing)))
            rows = [
                {key: (row.get(headers[key]) or "").strip() for key in REQUIRED_COLUMNS}
                for row in reader
            ]
        if not rows:
            raise ValueError("mapping CSV has no rows")

        seen_benchmark: set[str] = set()
        seen_persisted: set[str] = set()
        errors: list[str] = []
        for line, row in enumerate(rows, start=2):
            benchmark_id = row["benchmark_job_id"]
            persisted_id = row["persisted_job_id"]
            if not benchmark_id or benchmark_id in seen_benchmark:
                errors.append(f"row {line}: missing or duplicate benchmark_job_id")
            seen_benchmark.add(benchmark_id)
            if not persisted_id.isdigit() or int(persisted_id) <= 0:
                errors.append(f"row {line}: persisted_job_id must be a positive integer job_postings.id")
            if persisted_id in seen_persisted:
                errors.append(f"row {line}: persisted_job_id maps more than one benchmark job")
            seen_persisted.add(persisted_id)
            if not row["persisted_job_title"] or not row["persisted_company"]:
                errors.append(f"row {line}: persisted title and company are required")
            if not row["persisted_url"].startswith("https://"):
                errors.append(f"row {line}: persisted_url must be HTTPS")
            if row["source_table"] != "job_postings":
                errors.append(f"row {line}: source_table must be job_postings")
            if not re.fullmatch(r"[0-9a-fA-F]{64}", row["source_export_sha256"]):
                errors.append(f"row {line}: source_export_sha256 must be a SHA-256 hex digest")
            try:
                timestamp = datetime.fromisoformat(row["verified_at_utc"].replace("Z", "+00:00"))
                if timestamp.tzinfo is None:
                    raise ValueError("timezone missing")
            except ValueError:
                errors.append(f"row {line}: verified_at_utc must be ISO-8601 with timezone")

        missing_ids = sorted(expected_ids - seen_benchmark)
        unexpected_ids = sorted(seen_benchmark - expected_ids)
        if missing_ids:
            errors.append(f"unmapped benchmark job IDs ({len(missing_ids)}): {', '.join(missing_ids[:10])}")
        if unexpected_ids:
            errors.append(f"unexpected benchmark job IDs ({len(unexpected_ids)}): {', '.join(unexpected_ids[:10])}")
        if len(rows) != len(expected_ids):
            errors.append(f"mapping row count {len(rows)} does not match unique benchmark jobs {len(expected_ids)}")

        if errors:
            raise ValueError("\n".join(errors))

        workbook_hash = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
        mapping_hash = hashlib.sha256(args.mapping.read_bytes()).hexdigest()
        export_hashes = sorted({row["source_export_sha256"].lower() for row in rows})
        if len(export_hashes) != 1:
            raise ValueError("all rows must reference the same source export SHA-256")

        manifest = {
            "status": "mapping_schema_validated_pending_independent_review",
            "benchmark_workbook_sha256": workbook_hash,
            "mapping_csv_sha256": mapping_hash,
            "source_export_sha256": export_hashes[0],
            "unique_benchmark_jobs": len(expected_ids),
            "mapping_rows": len(rows),
            "unique_persisted_job_ids": len(seen_persisted),
            "source_table": "job_postings",
            "production_authoritative": False,
            "note": "Schema validation does not prove database authenticity or ranking quality; review source export and mapped records before promotion.",
        }
        args.manifest.parent.mkdir(parents=True, exist_ok=True)
        args.manifest.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(manifest, indent=2))
        return 0
    except (OSError, ValueError, KeyError) as exc:
        raise SystemExit(f"JOB MAPPING VALIDATION FAILED: {exc}") from exc


if __name__ == "__main__":
    raise SystemExit(main())
